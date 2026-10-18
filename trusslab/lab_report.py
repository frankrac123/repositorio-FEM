"""Portable Python report exporters for the desktop application."""
from pathlib import Path
import csv,json,html
from matplotlib.figure import Figure
from matplotlib.backends.backend_agg import FigureCanvasAgg
from .lab_plots import PLOT_TYPES,draw_plot
from .laboratory import winner_model
from .core import Model
from .analysis import design_cases


def report_text(campaign):
    model=campaign['model'];document=Model.from_dict(model);config=campaign['configuration'];lines=[
        '# Informe de comparación estructural','',f'Modelo: {model["name"]}',
        f'Huella: {campaign["model_fingerprint"]}',f'Estado: {campaign["status"]}',
        f'Configuración: {len(config["seeds"])} corridas por método; presupuesto {config["budget"]} por corrida; motor {config["engine"]}.','',
        '## Problema y formulación',
        f'Áreas {model["area_min"]} a {model["area_max"]} in². Propiedades por barra: '+ '; '.join(f'B{b.id}: E={document.member_material(b).elastic_modulus:g} ksi, gamma={document.member_material(b).density:g} lb/in³' for b in document.members),
        f'Esfuerzos admisibles ±{model["stress_limit"]} ksi; desplazamientos ±{model["displacement_limit"]} in por componente libre.',
        'Minimizar W = suma(gamma A L), en lb. Rigidez axial EA/L; solución Kff uf = Ff y reacciones Ku - F. El peso propio solicitado se concentra por mitades en los nodos extremos para conservar la formulación axial.',
        'Patrones: '+json.dumps(model.get('load_patterns',[]) or [{'id':'LC1','self_weight':0,'nodal_loads':{str(n['id']):n['load'] for n in model['nodes']}}],ensure_ascii=False),
        'Casos verificados en cada candidato: '+', '.join(design_cases(document)),
        'Tolerancia de factibilidad: 1e-6 en las relaciones normalizadas. Estabilidad: condición <= 1e13. Equilibrio: 1e-6 kip y 1e-3 kip in.','',
        '## Métodos y configuración experimental',
        'ASAM usa exploración uniforme, radio geométrico y aceptación logística. Las direcciones simétricas, reflexión de límites, normalización y tratamiento de restricciones son decisiones del proyecto.',
        'DE utiliza SciPy con restricciones explícitas y sin pulido oculto. SLSQP es determinista con varios arranques obtenidos de la misma población común; las semillas cambian los arranques.',
        'Los métodos que incluyen +SLSQP contabilizan el refinamiento dentro del presupuesto y se presentan por separado de los puros.',
        'Cada método recibe la misma población inicial por semilla. Se registran análisis nuevos, auditorías FEM del motor SAP, aciertos de caché, reinicios y gradientes.',
        'La verificación externa posterior no alimenta la búsqueda y tiene costes separados. No se introducen diseños publicados como puntos de arranque.',
        'Parámetros inmutables: '+json.dumps(config['parameters'],ensure_ascii=False),
        'Semillas: '+', '.join(map(str,config['seeds'])),
        'Versiones: '+json.dumps(campaign['software'],ensure_ascii=False),'',
        '## Resultados descriptivos',
        '| Método | n factibles/completadas | Mejor (lb) | Media (lb) | Mediana (lb) | Desviación muestral (lb) |',
        '|---|---:|---:|---:|---:|---:|']
    fmt=lambda x:'No disponible' if x is None else f'{x:.6f}'
    for s in campaign['statistics']:
        lines.append(f'| {s["algorithm"]} | {s["n_valid"]}/{s["n_completed"]} | {fmt(s["best_lb"])} | {fmt(s["mean_lb"])} | {fmt(s["median_lb"])} | {fmt(s["sample_std_lb"])} |')
    lines+=['','Tasa de factibilidad = corridas completas con diseño factible / corridas completas, incluyendo fallos y excluyendo canceladas. La desviación es muestral (n-1); no se define con una sola muestra.',
        'Con menos de 20 corridas factibles por método, la muestra es limitada. Estas estadísticas no prueban diferencias significativas ni superioridad universal.','',
        '## Validación SAP2000']
    for r in campaign['runs']:
        v=r['sap_validation'];lines.append(f'{r["algorithm"]}, semilla {r["seed"]}: {v["status"]}'+(' - '+v['error'] if v.get('error') else ''))
    lines+=['','## Restricciones activas','Se consideran próximas al límite con tolerancia relativa 1e-4; esto no modifica la tolerancia de factibilidad 1e-6.']
    for r in campaign['runs']:
        if r.get('result'):lines.append(f'{r["algorithm"]}, semilla {r["seed"]}: '+', '.join(r.get('active_constraints',[])))
    lines+=['Cumplimiento estricto exige relaciones <= 1. Los diseños aceptados con relaciones entre 1 y 1+1e-6 cumplen solo dentro de la tolerancia; cada resultado registra strict_feasible y feasibility_tolerance.', 'Coste de validación externa: '+json.dumps(campaign['validation_cost']),
        'Los SDB archivados son copias del modelo guardado del candidato verificado; los resultados numéricos y la comparación se conservan en JSON. Al abrir una copia aislada en SAP puede ser necesario volver a analizar para reconstruir sus archivos auxiliares.','',
        '## Discusión y conclusiones']
    candidates=[s for s in campaign['statistics'] if s['best_lb'] is not None]
    if candidates:
        best=min(candidates,key=lambda s:s['best_lb']);lines.append(f'El menor peso observado en esta campaña fue {best["best_lb"]:.6f} lb con {best["algorithm"]}. El criterio es descriptivo y depende del presupuesto y los arranques documentados.')
    else:lines.append('No hubo diseños factibles completos; no se declara un algoritmo ganador.')
    lines+=['No se sustituye una solución infactible o un análisis fallido por valores estructurales inventados. Los pendientes SAP no se presentan como validados.',
        'Los tiempos incluyen pausas y dependen del equipo y de la carga del sistema. Los métodos pueden finalizar antes del límite; se comparan también sus consumos reales.',
        'Referencias del Caso I: 5060.855730008 lb (diseño anterior), y aproximadamente 5060.877 lb (artículo). Solo son comparables si geometría, cargas, material y límites coinciden con el benchmark de 25 ksi; no son objetivos impuestos.','',
        '## Referencias',
        'Millán-Páramo, Matoski y Mazer (2017). DOI: https://doi.org/10.18845/tm.v30i2.3209',
        'SciPy: https://docs.scipy.org/doc/scipy/reference/generated/scipy.optimize.differential_evolution.html',
        'SciPy: https://docs.scipy.org/doc/scipy/reference/optimize.minimize-slsqp.html','']
    return '\n'.join(lines)


def export_report(campaign,directory,*,formats=('png','svg','pdf'),excel=True):
    """Export only recorded results. Never runs an optimizer or a SAP analysis."""
    directory=Path(directory);directory.mkdir(parents=True,exist_ok=True);histories={};figures=[];paths=[]
    text=report_text(campaign);(directory/'informe.md').write_text(text,encoding='utf-8')
    for kind in PLOT_TYPES:
        figure=Figure(figsize=(9,5.4));FigureCanvasAgg(figure);draw_plot(figure,campaign,kind,histories=histories)
        for extension in formats:
            path=directory/f'{kind}.{extension}';figure.savefig(path,dpi=170,bbox_inches='tight');paths.append(str(path))
        if 'png' not in formats:figure.savefig(directory/f'{kind}.png',dpi=170,bbox_inches='tight')
        figures.append(directory/f'{kind}.png');figure.clear()
    # A human-readable report and the structured campaign travel together.
    payload={k:v for k,v in campaign.items() if not k.startswith('_')}
    (directory/'campaign.json').write_text(json.dumps(payload,indent=2,ensure_ascii=False,allow_nan=False),encoding='utf-8')
    export_pdf(campaign,text,figures,directory/'informe.pdf');paths.append(str(directory/'informe.pdf'))
    if excel:export_excel(campaign,directory/'resultados.xlsx');paths.append(str(directory/'resultados.xlsx'))
    return paths


def export_pdf(campaign,text,figures,path):
    from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle,Image,PageBreak
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    import matplotlib
    fonts=Path(matplotlib.get_data_path())/'fonts'/'ttf'
    for name,file in [('LabSans','DejaVuSans.ttf'),('LabSansBold','DejaVuSans-Bold.ttf')]:
        if name not in pdfmetrics.getRegisteredFontNames():pdfmetrics.registerFont(TTFont(name,str(fonts/file)))
    styles=getSampleStyleSheet()
    for style in styles.byName.values():style.fontName='LabSans';style.fontSize=9;style.leading=13
    styles['Title'].fontName='LabSansBold';styles['Title'].fontSize=20;styles['Title'].leading=25
    styles['Heading1'].fontName='LabSansBold';styles['Heading1'].fontSize=12;styles['Heading1'].leading=17
    doc=SimpleDocTemplate(str(path),pagesize=A4,rightMargin=43,leftMargin=43,topMargin=48,bottomMargin=43)
    story=[]
    def para(value,style='BodyText'):return Paragraph(html.escape(str(value)).replace('\n','<br/>'),styles[style])
    for line in text.splitlines():
        if line.startswith('# '):story.append(para(line[2:],'Title'))
        elif line.startswith('## '):story.append(Spacer(1,10));story.append(para(line[3:],'Heading1'))
        elif line.startswith('|'):continue  # Render structured table below at readable widths.
        elif line:story.append(para(line));story.append(Spacer(1,5))
    story.append(PageBreak());story.append(para('Geometría, cargas y apoyos','Heading1'))
    rows=[['Nodo','X (in)','Y (in)','Z (in)','Restricciones','Fx/Fy/Fz (kip)']]
    document=Model.from_dict(campaign['model']);forces=document.nodal_loads()
    story.append(para('Cargas mostradas: caso '+document.active_case+'. Los patrones y factores completos constan en la definición del problema.'))
    for n in campaign['model']['nodes']:
        rows.append([n['id'],*n['xyz'],''.join(axis for axis,fixed in zip('XYZ',document.restraints(document.node(n['id']))) if fixed),'/'.join(f'{v:g}' for v in forces[str(n['id'])])])
    def table(rows,widths):
        t=Table([[para(v) for v in row] for row in rows],colWidths=widths,repeatRows=1,hAlign='LEFT')
        t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#dceafa')),('VALIGN',(0,0),(-1,-1),'TOP'),
            ('GRID',(0,0),(-1,-1),.4,colors.HexColor('#cad7e7')),('LEFTPADDING',(0,0),(-1,-1),5),('RIGHTPADDING',(0,0),(-1,-1),5),('TOPPADDING',(0,0),(-1,-1),6),('BOTTOMPADDING',(0,0),(-1,-1),6)]))
        story.append(t);story.append(Spacer(1,12))
    table(rows,[40,60,60,60,100,100])
    story.append(para('Estadísticas descriptivas','Heading1'));rows=[['Método','n válido','Mejor (lb)','Media (lb)','Desv. (lb)']]
    for s in campaign['statistics']:
        rows.append([s['algorithm'],f'{s["n_valid"]}/{s["n_completed"]}',*['—' if s[k] is None else f'{s[k]:.6f}' for k in ['best_lb','mean_lb','sample_std_lb']]])
    table(rows,[110,50,90,90,80])
    story.append(para('Áreas del mejor diseño factible por método','Heading1'))
    for method in campaign['configuration']['algorithms']:
        valid=[r for r in campaign['runs'] if r['algorithm']==method and r['status']=='feasible']
        if valid:
            r=min(valid,key=lambda r:r['result']['weight_lb']);story.append(para(f'{method}, semilla {r["seed"]}, caso mostrado {r["result"].get("analysis_case","LC1")}, SAP {r["sap_validation"]["status"]}'))
            table([['Barra','Inicio','Final','A (in²)','Axial (kip)','Esfuerzo (ksi)']]+[
                [b['id'],b['start'],b['end'],f'{a:.6f}',f'{v["force"]:.6f}',f'{v["stress"]:.6f}']
                for b,a,v in zip(campaign['model']['members'],r['areas_in2'],r['result']['members'])],[45,45,45,85,100,100])
    for figure in figures:
        story.append(PageBreak());story.append(para(PLOT_TYPES[figure.stem],'Heading1'));story.append(Image(str(figure),width=500,height=300,kind='proportional'))
    def footer(canvas,document):
        canvas.saveState();canvas.setFont('Helvetica',8);canvas.setFillColor(colors.HexColor('#526783'))
        canvas.drawString(43,24,'TrussLab3D | Datos de campaña registrados');canvas.drawRightString(A4[0]-43,24,str(document.page));canvas.restoreState()
    doc.build(story,onFirstPage=footer,onLaterPages=footer)


def export_excel(campaign,path):
    # Part of the portable Python application, not a Codex runtime dependency.
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment
    from openpyxl.utils import get_column_letter
    workbook=Workbook();workbook.remove(workbook.active)
    def sheet(name,headers,rows):
        ws=workbook.create_sheet(name);ws.append(headers)
        for row in rows:ws.append(row)
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='244C7D')
        for row in ws.iter_rows(min_row=2):
            for c in row:
                if isinstance(c.value,float):c.number_format='0.000000'
                c.alignment=Alignment(vertical='top')
        for i,header in enumerate(headers,1):ws.column_dimensions[get_column_letter(i)].width=min(44,max(16,len(header)+3))
        return ws
    columns=['algorithm','n_valid','n_completed','best_lb','worst_lb','mean_lb','median_lb','sample_std_lb','feasibility_rate','mean_seconds','mean_evaluations_to_best','n_failed_runs']
    sheet('Resumen',columns,[[s[k] for k in columns] for s in campaign['statistics']])
    headers=['Método','Semilla','Estado','Peso (lb)','Análisis nuevos','Caché','Tiempo (s)','SAP']
    sheet('Corridas',headers,[[r['algorithm'],r['seed'],r['status'],r.get('result',{}).get('weight_lb'),r['metrics']['evaluations'],r['metrics']['cache_hits'],r['metrics']['elapsed_seconds'],r['sap_validation']['status']] for r in campaign['runs']])
    sheet('Áreas', ['Método','Semilla','Barra','Área (in²)'],[[r['algorithm'],r['seed'],b['id'],a] for r in campaign['runs'] if r.get('areas_in2') for b,a in zip(campaign['model']['members'],r['areas_in2'])])
    sheet('Esfuerzos',['Método','Semilla','Caso','Barra','Axial (kip)','Esfuerzo (ksi)','Utilización'],[[r['algorithm'],r['seed'],case,v['id'],v['force'],v['stress'],v['ratio']] for r in campaign['runs'] if r.get('result') for case,result in (r['result'].get('case_results') or {r['result'].get('analysis_case','LC1'):r['result']}).items() for v in result['members']])
    sheet('Desplazamientos',['Método','Semilla','Caso','Nodo','ux (in)','uy (in)','uz (in)'],[[r['algorithm'],r['seed'],case,node,*u] for r in campaign['runs'] if r.get('result') for case,result in (r['result'].get('case_results') or {r['result'].get('analysis_case','LC1'):r['result']}).items() for node,u in result['displacements'].items()])
    sheet('Restricciones',['Método','Semilla','Restricción próxima al límite','Tolerancia relativa'],[[r['algorithm'],r['seed'],v,1e-4] for r in campaign['runs'] for v in r.get('active_constraints',[])])
    sheet('Configuración',['Campo','Valor'],[['Huella',campaign['model_fingerprint']],['Configuración',json.dumps(campaign['configuration'],ensure_ascii=False)],['Versiones',json.dumps(campaign['software'],ensure_ascii=False)],['Definición tasa','Factibles completas / completas; canceladas excluidas'],['Desviación','Muestral, n-1; n<2 no definida'],['Limitación','Menos de 20 corridas no demuestra superioridad estadística']])
    workbook.save(path)

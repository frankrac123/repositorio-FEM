import json,shutil
from pathlib import Path
import pytest
from test_laboratory import configuration
from trusslab.core import benchmark
from trusslab.fem import solve
from trusslab.laboratory import run_campaign,validate_campaign,winner_model


@pytest.mark.parametrize('method',['ASAM+SLSQP','DE+SLSQP'])
def test_hybrids_keep_identity_and_refinement_inside_total_budget(tmp_path,method):
    c=run_campaign(benchmark(),configuration([method],budget=360),tmp_path);r=c['runs'][0]
    assert r['algorithm']==method and r['search']['hybrid'] and r['metrics']['evaluations']<=360
    assert r['result']['weight_lb']<=r['search']['global_weight_lb']
    assert solve(winner_model(c,r))['feasible']
    assert any(s['phase']=='local_refinement' for s in r['search']['starts'])


def test_validation_error_retains_fem_pending_without_fallback(tmp_path):
    c=run_campaign(benchmark(),configuration(['ASAM']),tmp_path);before=c['runs'][0]['result']['weight_lb']
    class Broken:
        def __init__(self,path):pass
        def __enter__(self):return self
        def __exit__(self,*args):pass
        def evaluate(self,model):raise RuntimeError('SAP offline')
    validate_campaign(c,c['_directory'],batch_factory=Broken)
    assert c['runs'][0]['result']['weight_lb']==before and c['runs'][0]['sap_validation']['status']=='pending'
    assert 'SAP offline' in c['validation_error'] and c['validation_cost']['sap_analyses']==0


def test_validation_selects_one_best_per_method_with_one_owned_session(tmp_path):
    c=run_campaign(benchmark(),configuration(['ASAM'],seeds=[42,7]),tmp_path);created=[]
    class Batch:
        def __init__(self,path):created.append(self);self.calls=0;self.path=tmp_path/'dummy.sdb';self.path.write_bytes(b'UNIT TEST DOUBLE')
        def __enter__(self):return self
        def __exit__(self,*args):self.closed=True
        def evaluate(self,model):
            from trusslab.validation import validation_report
            self.calls+=1;r=solve(model);r.update(engine='SAP2000 TEST DOUBLE',sap_model=str(self.path),model_fingerprint=model.fingerprint(),analysis_case_status={'CASE_I':4})
            r['validation']=validation_report(model,r,r);return r
    validate_campaign(c,c['_directory'],batch_factory=Batch)
    assert len(created)==1 and created[0].calls==1 and created[0].closed
    verified=[r for r in c['runs'] if r['sap_validation']['status']=='verified'];assert len(verified)==1
    assert verified[0]['result']['weight_lb']==min(r['result']['weight_lb'] for r in c['runs'])


def test_report_exports_real_campaign_values_and_all_graph_formats(tmp_path):
    pytest.importorskip('reportlab');pytest.importorskip('openpyxl');pytest.importorskip('pypdf')
    from trusslab.lab_report import export_report
    from openpyxl import load_workbook
    from pypdf import PdfReader
    c=run_campaign(benchmark(),configuration(),tmp_path);out=tmp_path/'report';paths=export_report(c,out)
    for name in ['convergence','boxplot','time','evaluations','histogram','areas','feasibility']:
        assert all((out/f'{name}.{ext}').stat().st_size>100 for ext in ['png','svg','pdf'])
    workbook=load_workbook(out/'resultados.xlsx',data_only=True)
    assert workbook['Corridas'].cell(2,4).value==pytest.approx(c['runs'][0]['result']['weight_lb'])
    assert len(workbook['Áreas']['A'])==1+10*len(c['runs'])
    assert len(PdfReader(out/'informe.pdf').pages)>=8
    assert (out/'informe.md').is_file() and c['validation_cost']['sap_analyses']==0
